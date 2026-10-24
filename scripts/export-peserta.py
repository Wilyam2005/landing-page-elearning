from datetime import date, datetime
from pathlib import Path
import json
import re
from collections import Counter

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / "data history peserta  paket pkbm"
OUT = ROOT / "data" / "peserta.json"

ROMAN = {
    "VI": "6",
    "VII": "7",
    "VIII": "8",
    "IX": "9",
    "X": "10",
    "XI": "11",
    "XII": "12",
}


def cell(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    if s.upper() in ("NULL", "NONE", "NAN"):
        return ""
    return s


def looks_like_header(row):
    joined = " ".join(row).upper()
    return "NAMA" in joined and ("JK" in joined or "NISN" in joined)


def looks_like_title(row):
    j = " ".join(row).upper()
    return (
        "PESERTA DIDIK" in j
        or "DAFTAR" in j
        or "PKBM" in j
        or "TANGGAL UNDUH" in j
        or "KECAMATAN" in j
    )


def normalize_kelas(raw):
    if not raw:
        return "", "", "", ""
    s = raw.strip().upper().replace(".", "").replace("KELAS", "KLS")
    s = re.sub(r"\s+", " ", s)
    tingkat = ""
    rombel = ""
    m_ab = re.search(r"\b(XII|XI|X|IX|VIII|VII|VI)\s+([A-D])\b", s)
    m_rom = re.search(r"\b(XII|XI|IX|X|VIII|VII|VI)\b", s)
    m_num = re.search(r"(\d{1,2})", s)
    if m_ab:
        tingkat = ROMAN[m_ab.group(1)]
        rombel = m_ab.group(2)
    elif m_rom:
        tingkat = ROMAN[m_rom.group(1)]
        rest = s[m_rom.end() :].strip()
        if rest in ("A", "B", "C", "D"):
            rombel = rest
    elif m_num:
        tingkat = str(int(m_num.group(1)))
        rest = s[m_num.end() :].strip()
        if rest in ("A", "B", "C", "D"):
            rombel = rest
    label = f"Kelas {tingkat}" if tingkat else raw.strip()
    if rombel:
        label = f"Kelas {tingkat} {rombel}"
    jenjang = ""
    if tingkat:
        t = int(tingkat)
        if t <= 6:
            jenjang = "Paket A (SD)"
        elif t <= 9:
            jenjang = "Paket B (SMP)"
        else:
            jenjang = "Paket C (SMA)"
    return label, tingkat, jenjang, rombel


def parse_age(tgl, ref_year=2026):
    if not tgl or len(tgl) < 4:
        return None
    try:
        y = int(tgl[:4])
        if 1900 < y < 2027:
            return ref_year - y
    except ValueError:
        return None
    return None


def main():
    all_raw = []
    for f in sorted(FOLDER.glob("*.xlsx")):
        wb = load_workbook(f, data_only=True)
        year = "unknown"
        m = re.search(r"(\d{4})(\d{4})", f.name)
        if m:
            year = f"{m.group(1)}/{m.group(2)}"
        for name in wb.sheetnames:
            ws = wb[name]
            for r in range(1, ws.max_row + 1):
                vals = [cell(ws.cell(r, c).value) for c in range(1, 10)]
                if not any(vals):
                    continue
                if looks_like_header(vals) or looks_like_title(vals):
                    continue
                no, nama, jk, nisn, tempat, tgl, nik, alamat, kls = (vals + [""] * 9)[:9]
                if not nama or nama.upper() in ("NO", "NAMA"):
                    continue
                if not any([jk, nisn, tempat, tgl, nik, alamat, kls]) and (
                    nama.isdigit() or len(nama) < 2
                ):
                    continue
                all_raw.append(
                    {
                        "tahunAjaran": year,
                        "sumber": name,
                        "file": f.name,
                        "noAsli": no,
                        "nama": nama,
                        "jk": jk.upper()[:1] if jk else "",
                        "nisn": re.sub(r"\s+", "", nisn),
                        "tempatLahir": tempat,
                        "tanggalLahir": tgl,
                        "nik": re.sub(r"\s+", "", nik),
                        "alamat": alamat,
                        "kelasAsli": kls,
                    }
                )

    kelas_map = {}
    for r in all_raw:
        if r["tahunAjaran"] == "2026/2027" and r["sumber"] == "Sheet1" and r["kelasAsli"]:
            keys = []
            if r["nik"]:
                keys.append("nik:" + r["nik"])
            if r["nisn"]:
                keys.append("nisn:" + r["nisn"])
            keys.append("nama:" + r["nama"].upper())
            for k in keys:
                kelas_map[k] = r["kelasAsli"]

    records = []
    for r in all_raw:
        if r["tahunAjaran"] == "2026/2027" and r["sumber"] == "Sheet1":
            continue
        kelas = r["kelasAsli"]
        if not kelas and r["tahunAjaran"] == "2026/2027":
            for k in [
                ("nik:" + r["nik"]) if r["nik"] else None,
                ("nisn:" + r["nisn"]) if r["nisn"] else None,
                "nama:" + r["nama"].upper(),
            ]:
                if k and k in kelas_map:
                    kelas = kelas_map[k]
                    break
        label, tingkat, jenjang, rombel = normalize_kelas(kelas)
        records.append(
            {
                **{k: v for k, v in r.items() if k != "kelasAsli"},
                "kelasAsli": kelas or r["kelasAsli"],
                "kelas": label or "Belum terisi",
                "tingkat": tingkat,
                "jenjang": jenjang or "Belum terisi",
                "rombel": rombel,
                "tahunLahir": r["tanggalLahir"][:4]
                if r["tanggalLahir"] and len(r["tanggalLahir"]) >= 4
                else "",
                "usiaPerkiraan": parse_age(r["tanggalLahir"]),
            }
        )

    payload = {
        "sekolah": "PKBM Assyuro Lombok Creative (ALC)",
        "lokasi": "Kecamatan Masbagik, Kabupaten Lombok Timur, Nusa Tenggara Barat",
        "generatedAt": "2026-10-24",
        "tahunAjaran": sorted(set(r["tahunAjaran"] for r in records)),
        "total": len(records),
        "fields": [
            "tahunAjaran",
            "sumber",
            "file",
            "noAsli",
            "nama",
            "jk",
            "nisn",
            "tempatLahir",
            "tanggalLahir",
            "nik",
            "alamat",
            "kelasAsli",
            "kelas",
            "tingkat",
            "jenjang",
            "rombel",
            "tahunLahir",
            "usiaPerkiraan",
        ],
        "peserta": records,
    }
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    print("written", OUT, "bytes", OUT.stat().st_size)
    print("total", len(records))
    print("by year", Counter(r["tahunAjaran"] for r in records))
    print("by jenjang", Counter(r["jenjang"] for r in records))
    print("by jk", Counter(r["jk"] for r in records))
    print("kelas", sorted(set(r["kelas"] for r in records)))
    print("missing kelas", sum(1 for r in records if r["kelas"] == "Belum terisi"))


if __name__ == "__main__":
    main()
